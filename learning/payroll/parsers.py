import io
import csv
import logging
from decimal import Decimal
from typing import Dict, List, Any, Tuple
from openpyxl import load_workbook
from .security import sanitize_spreadsheet_value

logger = logging.getLogger('payroll.parsers')


def _find_column_index(headers: List[str], candidate_names: List[str]) -> int:
    """
    Finds the 0-indexed column matching any candidate header name (case-insensitive).
    """
    cleaned_headers = [h.strip().lower().replace('_', ' ') for h in headers]
    for candidate in candidate_names:
        c_clean = candidate.lower().replace('_', ' ')
        for idx, h in enumerate(cleaned_headers):
            if c_clean in h:
                return idx
    return -1


def parse_session_file(uploaded_file) -> Tuple[Dict[str, Dict[str, Any]], List[str]]:
    """
    Parses TutorCruncher or Zoho CSV/Excel session logs.
    Security: Applies sanitization against CSV/formula injection on all text fields.
    Returns: (grouped_tutors_dict, error_log_list)
    """
    file_name = uploaded_file.name.lower()
    rows = []

    if file_name.endswith('.csv'):
        content = uploaded_file.read().decode('utf-8-sig', errors='replace')
        reader = csv.reader(io.StringIO(content))
        rows = list(reader)
    elif file_name.endswith('.xlsx'):
        uploaded_file.seek(0)
        wb = load_workbook(filename=io.BytesIO(uploaded_file.read()), data_only=True, read_only=True)
        sheet = wb.active
        for r in sheet.iter_rows(values_only=True):
            rows.append([str(c) if c is not None else "" for c in r])

    if not rows or len(rows) < 2:
        return {}, ["The uploaded file is empty or does not contain a header row."]

    headers = [str(c).strip() for c in rows[0]]
    
    # Header resolution with smart fallbacks
    tutor_name_idx = _find_column_index(headers, ['tutor', 'contractor', 'teacher', 'tutor name', 'staff'])
    tutor_email_idx = _find_column_index(headers, ['tutor email', 'contractor email', 'email', 'teacher email'])
    student_idx = _find_column_index(headers, ['student', 'student name', 'client', 'child', 'pupil'])
    duration_idx = _find_column_index(headers, ['duration', 'hours', 'units', 'duration hours', 'length'])
    level_idx = _find_column_index(headers, ['class', 'level', 'year', 'grade', 'service', 'class level'])
    date_idx = _find_column_index(headers, ['date', 'session date', 'start', 'start time'])
    tz_idx = _find_column_index(headers, ['timezone', 'tz', 'client timezone', 'country', 'location'])
    mode_idx = _find_column_index(headers, ['mode', 'delivery', 'lesson type', 'type'])

    if tutor_name_idx == -1 and tutor_email_idx == -1:
        return {}, ["Could not identify Tutor column in file. Expected 'Tutor' or 'Tutor Email'."]

    grouped_tutors: Dict[str, Dict[str, Any]] = {}
    errors: List[str] = []

    for row_num, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue  # Skip blank lines

        try:
            tutor_name = sanitize_spreadsheet_value(row[tutor_name_idx] if tutor_name_idx != -1 and tutor_name_idx < len(row) else "Unknown Tutor")
            tutor_email = sanitize_spreadsheet_value(row[tutor_email_idx] if tutor_email_idx != -1 and tutor_email_idx < len(row) else "").lower()
            
            tutor_key = tutor_email if tutor_email else tutor_name.lower()

            student_name = sanitize_spreadsheet_value(row[student_idx] if student_idx != -1 and student_idx < len(row) else "Student")
            date_val = sanitize_spreadsheet_value(row[date_idx] if date_idx != -1 and date_idx < len(row) else "")
            class_level = sanitize_spreadsheet_value(row[level_idx] if level_idx != -1 and level_idx < len(row) else "year 7")
            tz_val = sanitize_spreadsheet_value(row[tz_idx] if tz_idx != -1 and tz_idx < len(row) else "WAT")
            mode_val = sanitize_spreadsheet_value(row[mode_idx] if mode_idx != -1 and mode_idx < len(row) else "Online lessons")

            # Duration conversion
            raw_duration = row[duration_idx] if duration_idx != -1 and duration_idx < len(row) else "1.0"
            try:
                # Handle formats like '1.5' or '1:30'
                if ':' in str(raw_duration):
                    parts = str(raw_duration).split(':')
                    duration_hours = Decimal(parts[0]) + (Decimal(parts[1]) / Decimal(60))
                else:
                    duration_hours = Decimal(str(raw_duration).strip().replace('h', '').replace('hrs', ''))
            except Exception:
                duration_hours = Decimal('1.00')

            if tutor_key not in grouped_tutors:
                grouped_tutors[tutor_key] = {
                    'tutor_name': tutor_name,
                    'tutor_email': tutor_email,
                    'sessions': []
                }

            grouped_tutors[tutor_key]['sessions'].append({
                'date': date_val,
                'student_name': student_name,
                'class_level': class_level,
                'duration_hours': duration_hours,
                'timezone': tz_val,
                'delivery_mode': mode_val,
            })

        except Exception as e:
            errors.append(f"Row {row_num}: Error processing entry - {str(e)}")

    return grouped_tutors, errors
