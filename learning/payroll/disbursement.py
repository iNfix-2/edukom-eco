import io
import csv
from decimal import Decimal
from typing import List, Dict, Any, Tuple
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .models import PayrollRun, TutorPayrollItem, TutorBankAccount
from .security import sanitize_spreadsheet_value


def generate_paystack_transfer_batch(payroll_run: PayrollRun) -> Tuple[List[Dict[str, Any]], List[str]]:
    """
    Generates Paystack Bulk Transfer payload from an approved PayrollRun.
    Returns: (transfers_payload_list, warnings_list)
    """
    transfers = []
    warnings = []

    for item in payroll_run.items.all():
        if item.net_pay <= Decimal('0.00'):
            continue

        bank_acct = None
        if item.tutor:
            bank_acct = getattr(item.tutor, 'bank_account', None)

        if not bank_acct or not bank_acct.recipient_code:
            warnings.append(
                f"Tutor '{item.tutor_name}' ({item.tutor_email}) does not have a verified Paystack recipient code. Skipped from automated transfer."
            )
            continue

        amount_in_kobo = int(item.net_pay * 100)
        ref = f"PAY-{payroll_run.id.hex[:8]}-{item.id.hex[:8]}"

        transfers.append({
            'amount': amount_in_kobo,
            'recipient': bank_acct.recipient_code,
            'reference': ref,
            'reason': f"Edukom Payroll - {payroll_run.title}"[:100]
        })

    return transfers, warnings


def export_payroll_excel(payroll_run: PayrollRun) -> io.BytesIO:
    """
    Exports a professionally formatted, formula-injection sanitized Excel workbook
    for internal accounting and management records.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Payroll Summary"

    # Styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Deep navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title block
    ws.merge_cells('A1:I1')
    ws['A1'] = sanitize_spreadsheet_value(f"EDUKOM GROUP - {payroll_run.title.upper()}")
    ws['A1'].font = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    ws['A1'].alignment = Alignment(horizontal="left", vertical="center")

    ws['A2'] = "Period:"
    ws['B2'] = f"{payroll_run.period_start} to {payroll_run.period_end}"
    ws['A3'] = "Status:"
    ws['B3'] = payroll_run.get_status_display()
    ws['A2'].font = bold_font
    ws['A3'].font = bold_font

    headers = [
        "Tutor Name",
        "Email",
        "Bank Name",
        "Account Number",
        "Total Hours",
        f"Gross Pay ({payroll_run.currency})",
        f"Deductions ({payroll_run.currency})",
        f"Net Pay ({payroll_run.currency})",
        "Payment Status"
    ]

    header_row = 5
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Rows
    current_row = 6
    for item in payroll_run.items.all():
        bank_name = "N/A"
        acct_num = "N/A"
        if item.tutor and hasattr(item.tutor, 'bank_account'):
            bank_name = item.tutor.bank_account.bank_name
            acct_num = item.tutor.bank_account.masked_account_number

        row_data = [
            sanitize_spreadsheet_value(item.tutor_name),
            sanitize_spreadsheet_value(item.tutor_email),
            sanitize_spreadsheet_value(bank_name),
            sanitize_spreadsheet_value(acct_num),
            float(item.total_hours),
            float(item.calculated_gross_pay),
            float(item.deductions),
            float(item.net_pay),
            sanitize_spreadsheet_value(item.status),
        ]

        for col_num, val in enumerate(row_data, 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.value = val
            cell.font = regular_font
            cell.border = thin_border
            if col_num in (5, 6, 7, 8):
                cell.alignment = right_align
                if col_num != 5:
                    cell.number_format = '#,##0.00'

        current_row += 1

    # Totals Row
    ws.cell(row=current_row, column=1, value="TOTALS").font = bold_font
    ws.cell(row=current_row, column=5, value=float(payroll_run.total_hours)).font = bold_font
    ws.cell(row=current_row, column=6, value=float(payroll_run.total_gross)).font = bold_font
    ws.cell(row=current_row, column=8, value=float(payroll_run.total_net)).font = bold_font

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
